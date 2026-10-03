import base64
import csv
import io
import re

from odoo import _, fields, models
from odoo.exceptions import UserError


class CheckBookImportWizard(models.TransientModel):
    _name = "check.book.import.wizard"
    _description = "Import Checkbook Leaves"

    book_id = fields.Many2one("check.book", required=True, readonly=True)
    numbers = fields.Text(help="Leaf numbers separated by new lines, commas or spaces.")
    file = fields.Binary(help="CSV or Excel file; the first column holds the leaf numbers.")
    filename = fields.Char()

    def action_import(self):
        self.ensure_one()
        numbers = self._collected_numbers()
        if not numbers:
            raise UserError(_("Enter or upload at least one leaf number."))
        self.book_id._import_leaves(numbers)
        return {"type": "ir.actions.act_window_close"}

    def _collected_numbers(self):
        seen, ordered = set(), []
        for number in self._typed_numbers() + self._file_numbers():
            if number not in seen:
                seen.add(number)
                ordered.append(number)
        return ordered

    def _typed_numbers(self):
        return [token.strip() for token in re.split(r"[\s,;]+", self.numbers or "") if token.strip()]

    def _file_numbers(self):
        if not self.file:
            return []
        content = base64.b64decode(self.file)
        if (self.filename or "").lower().endswith(".csv"):
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            values = [row[0] for row in rows if row]
        else:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = [row[0] for row in sheet.iter_rows(values_only=True) if row and row[0] not in (None, "")]
        cleaned = [str(value).strip() for value in values]
        return [value for value in cleaned if value and not value.lower().startswith("number")]
