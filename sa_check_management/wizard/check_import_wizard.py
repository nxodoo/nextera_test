import base64
import csv
import io
from datetime import date, datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..models.check_import import IMPORT_STATES

COLUMNS = [
    "direction", "check_number", "partner", "amount", "currency", "issue_date", "due_date",
    "bank", "drawer_name", "drawer_account", "purpose", "state", "journal", "deposit_date", "reference",
]
REQUIRED = ("direction", "check_number", "partner", "amount", "issue_date", "due_date", "state")
PURPOSES = ("payment", "guarantee", "security", "advance")
MAX_ERRORS_SHOWN = 30


class CheckImportWizard(models.TransientModel):
    _name = "check.import.wizard"
    _description = "Import Legacy Checks"

    file = fields.Binary(string="File", help="Excel (.xlsx) or CSV with the template columns.")
    filename = fields.Char()
    post_entries = fields.Boolean(
        string="Post Opening Check Entries", default=True,
        help="Post the receipt/issue entry of each open payment check (Checks Receivable/Payable against "
             "the partner). Leave unticked if the opening balances already contain the checks.",
    )
    template_file = fields.Binary(readonly=True)
    result = fields.Text(readonly=True)

    def action_download_template(self):
        self.ensure_one()
        self.template_file = base64.b64encode(self._template_bytes())
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/?model={self._name}&id={self.id}&field=template_file"
                   f"&filename=checks_import_template.xlsx&download=true",
            "target": "self",
        }

    def action_validate(self):
        self.ensure_one()
        rows, errors = self._parse_and_check()
        self.result = "\n".join(errors[:MAX_ERRORS_SHOWN]) if errors else _("%s rows are ready to import.", len(rows))
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        rows, errors = self._parse_and_check()
        if errors:
            raise UserError("\n".join(errors[:MAX_ERRORS_SHOWN]))
        checks = self.env["check.check"]
        for row in rows:
            check = self.env["check.check"].with_context(default_check_type=row["direction"]).create(row["vals"])
            check._import_into_state(row["state"], row["journal"], row["deposit_date"], self.post_entries)
            checks |= check
        return {
            "type": "ir.actions.act_window",
            "name": _("Imported Checks"),
            "res_model": "check.check",
            "views": [[False, "list"], [False, "form"]],
            "domain": [("id", "in", checks.ids)],
            "context": {"create": False},
        }

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------
    def _read_rows(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Upload a file first."))
        content = base64.b64decode(self.file)
        if (self.filename or "").lower().endswith(".csv"):
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        header = [str(cell or "").strip().lower() for cell in values[0]]
        return [dict(zip(header, row)) for row in values[1:] if any(cell not in (None, "") for cell in row)]

    def _parse_and_check(self):
        rows, errors = [], []
        for index, raw in enumerate(self._read_rows(), start=2):
            try:
                rows.append(self._parse_row(raw))
            except UserError as exc:
                errors.append(_("Row %(row)s: %(error)s", row=index, error=exc.args[0]))
        if not rows and not errors:
            errors.append(_("The file has no data rows."))
        return rows, errors

    def _parse_row(self, raw):
        missing = [column for column in REQUIRED if not raw.get(column) and raw.get(column) != 0]
        if missing:
            raise UserError(_("missing %s", ", ".join(missing)))
        direction = str(raw["direction"]).strip().lower()
        if direction not in IMPORT_STATES:
            raise UserError(_("direction must be incoming or outgoing"))
        state = str(raw["state"]).strip().lower()
        if state not in IMPORT_STATES[direction]:
            raise UserError(_("state for %(direction)s checks must be one of %(states)s",
                              direction=direction, states=", ".join(IMPORT_STATES[direction])))
        purpose = str(raw.get("purpose") or "payment").strip().lower()
        if purpose not in PURPOSES:
            raise UserError(_("purpose must be one of %s", ", ".join(PURPOSES)))
        journal = self._find_journal(raw.get("journal"))
        if (direction == "outgoing" or state == "under_collection") and not journal:
            raise UserError(_("journal is required for outgoing and deposited checks"))
        vals = {
            "check_number": str(raw["check_number"]).strip(),
            "partner_id": self._find_partner(raw["partner"]).id,
            "amount": self._to_amount(raw["amount"]),
            "currency_id": self._find_currency(raw.get("currency")).id,
            "issue_date": self._to_date(raw["issue_date"]),
            "due_date": self._to_date(raw["due_date"]),
            "purpose": purpose,
            "reference": raw.get("reference") or False,
        }
        if direction == "incoming":
            vals.update({
                "bank_id": self._find_bank(raw.get("bank")).id,
                "drawer_account_number": str(raw.get("drawer_account") or "").strip(),
                "drawer_is_partner": not raw.get("drawer_name"),
                "drawer_name": raw.get("drawer_name") or False,
            })
        else:
            vals["journal_id"] = journal.id
        return {
            "direction": direction, "state": state, "vals": vals, "journal": journal,
            "deposit_date": self._to_date(raw["deposit_date"]) if raw.get("deposit_date") else False,
        }

    # ------------------------------------------------------------------
    # Lookups and conversions (one each)
    # ------------------------------------------------------------------
    def _find_partner(self, value):
        value = str(value).strip()
        Partner = self.env["res.partner"]
        partner = Partner.search([("ref", "=", value)], limit=2) or Partner.search([("name", "=ilike", value)], limit=2)
        if len(partner) != 1:
            raise UserError(_("partner %(value)s not found or ambiguous (use its internal reference)", value=value))
        return partner

    def _find_bank(self, value):
        if not value:
            raise UserError(_("bank is required for incoming checks"))
        value = str(value).strip()
        bank = self.env["res.bank"].search(["|", ("bic", "=", value), ("name", "=ilike", value)], limit=1)
        return bank or self.env["res.bank"].create({"name": value})

    def _find_journal(self, value):
        if not value:
            return self.env["account.journal"]
        value = str(value).strip()
        journal = self.env["account.journal"].search([
            ("type", "=", "bank"), ("company_id", "=", self.env.company.id),
            "|", ("code", "=", value), ("name", "=ilike", value),
        ], limit=1)
        if not journal:
            raise UserError(_("bank journal %s not found", value))
        return journal

    def _find_currency(self, value):
        if not value:
            return self.env.company.currency_id
        currency = self.env["res.currency"].with_context(active_test=False).search([("name", "=ilike", str(value).strip())], limit=1)
        if not currency:
            raise UserError(_("currency %s not found", value))
        return currency

    @api.model
    def _to_amount(self, value):
        try:
            amount = float(str(value).replace(",", ""))
        except ValueError:
            raise UserError(_("amount %s is not a number", value)) from None
        if amount <= 0:
            raise UserError(_("amount must be positive"))
        return amount

    @api.model
    def _to_date(self, value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
        raise UserError(_("date %s must look like 2026-10-31 or 31/10/2026", value))

    @api.model
    def _template_bytes(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Checks"
        sheet.append(COLUMNS)
        sheet.append(["incoming", "845621", "Customer reference or name", 250000, "EGP", "2026-09-22", "2026-10-30",
                      "CIB", "", "100023459", "payment", "received", "", "", ""])
        sheet.append(["outgoing", "100045", "Vendor reference or name", 50000, "EGP", "2026-09-25", "2026-11-15",
                      "", "", "", "payment", "delivered", "BNK1", "", ""])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
